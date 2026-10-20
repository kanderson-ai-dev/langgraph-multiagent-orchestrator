"""DOCX export — renders the final report markdown as a .docx download.

Minimal block renderer (headings, bullet/numbered lists, paragraphs,
bold/italic/code runs and real hyperlinks) built on ``python-docx`` — a
single pure-Python dependency already pulled in with ``lxml`` underneath.
Anything richer would need pandoc; that trade-off is deliberately not
taken.
"""

from __future__ import annotations

import io
import re

import docx.opc.constants
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

_MD_LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")
_MD_INLINE = re.compile(
    r"\*\*(.+?)\*\*"  # 1: bold
    r"|(?<!\w)\*([^*\n]+)\*(?!\w)"  # 2: *italic*
    r"|(?<!\w)_([^_\n]+)_(?!\w)"  # 3: _italic_
    r"|`([^`]+)`"  # 4: code
)
_MD_HEADING = re.compile(r"^(#{1,6})\s*")
_MD_ULIST = re.compile(r"^[-*]\s+")
_MD_OLIST = re.compile(r"^(\d+)\.\s+")
_MD_HR = re.compile(r"^(-{3,}|\*{3,})$")


def _add_hyperlink(paragraph: Paragraph, url: str, text: str) -> None:
    """Append a real, clickable hyperlink run (Word's blue+underline)."""
    part = paragraph.part
    r_id = part.relate_to(
        url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    rpr.append(style)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _write_styled(paragraph: Paragraph, text: str) -> None:
    """Emit plain/bold/italic/code runs for a link-free text span."""
    pos = 0
    for m in _MD_INLINE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        bold, star, usc, code = m.groups()
        if bold is not None:
            run = paragraph.add_run(bold)
            run.bold = True
        elif star is not None or usc is not None:
            run = paragraph.add_run(star if star is not None else usc)
            run.italic = True
        elif code is not None:
            run = paragraph.add_run(code)
            run.font.name = "Consolas"
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _write_runs(paragraph: Paragraph, text: str) -> None:
    """Split ``[label](url)`` links out, style the rest inline."""
    pos = 0
    for m in _MD_LINK.finditer(text):
        _write_styled(paragraph, text[pos : m.start()])
        _add_hyperlink(paragraph, m.group(2), m.group(1))
        pos = m.end()
    _write_styled(paragraph, text[pos:])


def markdown_to_docx(
    markdown: str, *, title: str = "Report", footer: str = ""
) -> bytes:
    """Convert markdown text to DOCX bytes (headings, lists, body text)."""
    doc = Document()
    doc.add_heading(title[:120], level=0)

    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped or _MD_HR.match(stripped):
            continue
        m = _MD_HEADING.match(stripped)
        if m:
            level = min(len(m.group(1)), 4)
            p = doc.add_heading("", level=level)
            _write_runs(p, _MD_HEADING.sub("", stripped))
            continue
        m = _MD_OLIST.match(stripped)
        if m:
            # Literal numbers + hanging indent: python-docx's "List Number"
            # style shares one counter across the document, so separate
            # lists would keep counting instead of restarting.
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            _write_runs(p, f"{m.group(1)}. {_MD_OLIST.sub('', stripped)}")
            continue
        if _MD_ULIST.match(stripped):
            p = doc.add_paragraph(style="List Bullet")
            _write_runs(p, _MD_ULIST.sub("", stripped))
            continue
        p = doc.add_paragraph()
        _write_runs(p, stripped)

    if footer:
        p = doc.add_paragraph()
        run = p.add_run(footer)
        run.italic = True
        run.font.size = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
